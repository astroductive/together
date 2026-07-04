# -*- coding: utf-8 -*-
"""Together — project roll-up banner, 16.5 x 41 in (same print format as the
faculty's reference banner), rendered at 150 dpi and packaged as a one-slide
PPTX + PNG. Together design language throughout."""
import os
import sys
from PIL import Image, ImageDraw, ImageFont

HERE = os.path.dirname(os.path.abspath(__file__))
PRES = os.path.dirname(HERE)
sys.path.insert(0, PRES)
from deck_engine import (F, text_w, wrap, fit_text, paste_fit, ASSETS,
                         SERIF_B, SERIF_I, SANS_R, SANS_B)

DPI = 150
BW, BH = int(16.5 * DPI), int(41 * DPI)          # 2475 x 6150
M = 120

TEAL = "#0d8f83"
TEAL_DEEP = "#0a6b62"
TEAL_DARKER = "#075a52"
SAND = "#b97f3e"
SAND_L = "#e2b483"
INK = "#11233a"
BODY = "#34343d"
MUTED = "#545868"
FAINT = "#969dab"
HAIR = "#e4e7ee"
CARD = "#f4f7f7"
CREAM = "#f7f6f2"
WHITE = "#ffffff"


def A(name):
    import glob
    return glob.glob(os.path.join(ASSETS, name + ".*"))[0]


img = Image.new("RGB", (BW, BH), WHITE)
d = ImageDraw.Draw(img)


def vgrad(box, c1, c2):
    x0, y0, x1, y1 = box
    r1, g1, b1 = tuple(int(c1[i:i+2], 16) for i in (1, 3, 5))
    r2, g2, b2 = tuple(int(c2[i:i+2], 16) for i in (1, 3, 5))
    for y in range(y0, y1):
        t = (y - y0) / max(1, y1 - y0)
        d.line([(x0, y), (x1, y)],
               fill=(int(r1 + (r2 - r1) * t), int(g1 + (g2 - g1) * t),
                     int(b1 + (b2 - b1) * t)))


def logo_at(path, cx, cy, height, alpha=1.0):
    pic = Image.open(path).convert("RGBA")
    r = height / pic.height
    pic = pic.resize((int(pic.width * r), height), Image.LANCZOS)
    if alpha < 1.0:
        a = pic.split()[3].point(lambda v: int(v * alpha))
        pic.putalpha(a)
    img.paste(pic, (int(cx - pic.width / 2), int(cy - pic.height / 2)), pic)


def section_header(y, kicker, title):
    d.rectangle([M, y + 6, M + 16, y + 58], fill=SAND)
    d.text((M + 44, y), kicker.upper(), font=F(SANS_B, 34), fill=TEAL)
    d.text((M + 44, y + 46), title, font=F(SERIF_B, 72), fill=INK)
    d.line([(M, y + 150), (BW - M, y + 150)], fill=HAIR, width=3)
    return y + 178


def caption(box_bottom_center, text):
    cx, y = box_bottom_center
    f = F(SANS_R, 27)
    d.text((cx - text_w(f, text) / 2, y), text, font=f, fill=FAINT)


def arrow(x0, x1, cy, color=FAINT):
    d.line([(x0, cy), (x1 - 16, cy)], fill=color, width=5)
    d.polygon([(x1, cy), (x1 - 26, cy - 13), (x1 - 26, cy + 13)], fill=color)


# ═════════════════════ 1. top strip: FUE + department ═════════════════════
d.rectangle([0, 0, BW, 16], fill=SAND)
logo_at(A("fue_lockup"), 320, 110, 130)
fit_text(d, (620, 60, BW - M, 105),
         "Future University in Egypt — Faculty of Engineering and Technology",
         SANS_B, 34, 24, INK, 0, align="right")
fit_text(d, (620, 112, BW - M, 155),
         "Computers & Intelligent Systems Engineering — Graduation Project 2026",
         SANS_R, 28, 20, MUTED, 0, align="right")

# ═════════════════════ 2. hero ════════════════════════════════════════════
HERO_T, HERO_B = 200, 950
vgrad((0, HERO_T, BW, HERO_B), TEAL, TEAL_DARKER)
# oversized watermark hands, right side
logo_at(A("logo_white_hi"), BW - 320, HERO_B - 220, 620, alpha=0.14)
logo_at(A("logo_white_hi"), BW // 2, HERO_T + 150, 235)
fit_text(d, (M, HERO_T + 285, BW - M, HERO_T + 510), "Together",
         SERIF_B, 225, 150, WHITE, 0, align="center", line_gap=1.0)
fit_text(d, (M, HERO_T + 512, BW - M, HERO_T + 588),
         "An AI-Based Sign-Language Translator — ASL & ArSL",
         SERIF_B, 60, 40, "#eaf6f3", 0, align="center")
d.line([(BW // 2 - 140, HERO_T + 622), (BW // 2 + 140, HERO_T + 622)], fill=SAND_L, width=8)
fit_text(d, (M + 120, HERO_T + 650, BW - M - 120, HERO_T + 738),
         "Real-time, two-way translation between signed and spoken language — "
         "in the browser, with nothing but a webcam.",
         SERIF_I, 42, 30, "#d7efe9", 0, align="center", line_gap=1.25)

# ═════════════════════ 3. team band ═══════════════════════════════════════
TB_T, TB_H = HERO_B, 260
d.rectangle([0, TB_T, BW, TB_T + TB_H], fill=CREAM)
d.text((M, TB_T + 26), "PRESENTED BY", font=F(SANS_B, 30), fill=SAND)
names = [("Abdelfattah Moustafa", "20225889"), ("Michael Abdallah", "20213601"),
         ("Ahmed Mohamed Nagib", "20170423"), ("Ahmed Ashraf Shawareb", "20180097")]
cw = (BW - 2 * M - 3 * 40) // 4
for i, (nm, sid) in enumerate(names):
    x = M + i * (cw + 40)
    d.rounded_rectangle([x, TB_T + 74, x + cw, TB_T + 190], radius=18,
                        fill=WHITE, outline=HAIR, width=3)
    fit_text(d, (x + 20, TB_T + 96, x + cw - 20, TB_T + 140), nm,
             SANS_B, 34, 24, INK, 0, align="center")
    fit_text(d, (x + 20, TB_T + 142, x + cw - 20, TB_T + 180), sid,
             SANS_R, 30, 22, MUTED, 0, align="center")
sup = "Supervised by  Prof. Medhat Awadallah"
f = F(SANS_B, 33)
d.text(((BW - text_w(f, sup)) / 2, TB_T + 208), sup, font=f, fill=BODY)

# ═════════════════════ 4. motivation & objectives ═════════════════════════
y = section_header(TB_T + TB_H + 56, "The problem", "Why Together exists")
fit_text(d, (M, y, BW - M, y + 160),
         "Over 1.5 billion people live with hearing loss, yet everyday tools that "
         "translate between signed and spoken language remain one-directional, "
         "English-centric, and tied to special hardware. Together closes that gap for "
         "both American Sign Language and Egyptian Arabic Sign Language.",
         SANS_R, 40, 30, BODY, 0, line_gap=1.3)
y += 186
cards = [
    ("Bidirectional", "Sign → text & speech, and text & speech → a signing avatar — plus a live two-person meeting over WebRTC."),
    ("Bilingual", "ASL (250 signs) and Egyptian ArSL (20 signs) as first-class languages, each with its own trained model."),
    ("Hardware-free", "MediaPipe landmarks in the browser — no gloves, no depth cameras, no install, no GPU."),
]
cw = (BW - 2 * M - 2 * 44) // 3
for i, (t, b) in enumerate(cards):
    x = M + i * (cw + 44)
    d.rounded_rectangle([x, y, x + cw, y + 270], radius=20, fill=CARD,
                        outline=HAIR, width=3)
    d.rectangle([x, y, x + cw, y + 12], fill=TEAL)
    fit_text(d, (x + 34, y + 32, x + cw - 30, y + 90), t, SANS_B, 42, 30, TEAL_DEEP, 0)
    fit_text(d, (x + 34, y + 100, x + cw - 30, y + 256), b, SANS_R, 32, 24, BODY, 0, line_gap=1.26)
y += 270 + 56

# ═════════════════════ 5. system pipeline (drawn native) ══════════════════
y = section_header(y, "System design", "One architecture, four translation paths")
stages = [
    ("Webcam", "video in the browser,\nnothing installed"),
    ("MediaPipe Holistic", "543 landmarks\nper frame"),
    ("Recognition", "TFLite · 250 ASL\nPyTorch · 20 ArSL"),
    ("Gloss engine", "Topic-Comment ↔ SVO\n+ non-manual markers"),
    ("LLM sentence layer", "Gemini → local\nfallbacks"),
    ("Speech & avatar", "TTS out ·\nsigning avatar back"),
]
n = len(stages)
gap_a = 64
chip_w = (BW - 2 * M - (n - 1) * gap_a) // n
chip_h = 210
for i, (t, sub) in enumerate(stages):
    x = M + i * (chip_w + gap_a)
    d.rounded_rectangle([x, y, x + chip_w, y + chip_h], radius=18, fill=WHITE,
                        outline=HAIR, width=3)
    d.rectangle([x, y, x + chip_w, y + 10], fill=TEAL)
    d.ellipse([x + chip_w // 2 - 26, y - 26, x + chip_w // 2 + 26, y + 26], fill=SAND)
    nf = F(SANS_B, 30)
    d.text((x + chip_w // 2 - text_w(nf, str(i + 1)) / 2, y - 20), str(i + 1),
           font=nf, fill=WHITE)
    fit_text(d, (x + 18, y + 44, x + chip_w - 18, y + 96), t,
             SANS_B, 31, 22, INK, 0, align="center")
    for j, ln in enumerate(sub.split("\n")):
        lf = F(SANS_R, 26)
        d.text((x + chip_w // 2 - text_w(lf, ln) / 2, y + 108 + j * 38), ln,
               font=lf, fill=MUTED)
    if i < n - 1:
        arrow(x + chip_w + 8, x + chip_w + gap_a - 8, y + chip_h // 2, TEAL)
yy = y + chip_h + 46
paths = ["Sign → Text", "Sign → Speech", "Speech → Sign (avatar)", "Text → Sign (avatar)"]
pw = (BW - 2 * M - 3 * 44) // 4
for i, p in enumerate(paths):
    x = M + i * (pw + 44)
    d.rounded_rectangle([x, yy, x + pw, yy + 82], radius=41, fill=CARD,
                        outline=TEAL, width=3)
    fit_text(d, (x + 20, yy + 20, x + pw - 20, yy + 62), p, SANS_B, 32, 22,
             TEAL_DEEP, 0, align="center")
yy += 118
note = ("One FastAPI + Socket.IO server · PostgreSQL + pgvector semantic sign memory · "
        "WebRTC live meetings — every path runs on CPU only.")
nf = F(SANS_R, 28)
d.text(((BW - text_w(nf, note)) / 2, yy), note, font=nf, fill=FAINT)
y = yy + 60

# ═════════════════════ 6. the two models (cream band) ═════════════════════
band_t = y - 6
box_h = 1030
band_b = band_t + 178 + box_h + 48
d.rectangle([0, band_t, BW, band_b], fill=CREAM)
y = section_header(y, "Recognition", "Two models, sized to their data")
col_w = (BW - 2 * M - 60) // 2
for i, (title, fig, specs) in enumerate([
    ("ASL — 250 signs", "Figure4_3",
     ["Custom hybrid: Conv1D blocks + Transformer",
      "Input (60, 543, 3) MediaPipe landmarks → 708 features",
      "3-tower ensemble, exported to TFLite (CPU, XNNPACK)",
      "RAdam + Lookahead · cosine LR · 400 epochs · AWP"]),
    ("ArSL — 20 Egyptian signs", "Figure4_12",
     ["Compact CNN + 2-layer bidirectional GRU (64 units)",
      "Input (30, 177): 59 landmarks, shoulder-normalized",
      "INT8-quantized PyTorch — 4.2 ms per inference",
      "Adam · ReduceLROnPlateau · early stopping (12)"]),
]):
    x = M + i * (col_w + 60)
    d.rounded_rectangle([x, y, x + col_w, y + box_h], radius=22, fill=WHITE,
                        outline=HAIR, width=3)
    d.rectangle([x, y, x + col_w, y + 14], fill=TEAL if i == 0 else SAND)
    fit_text(d, (x + 36, y + 38, x + col_w - 30, y + 98), title,
             SERIF_B, 50, 36, INK, 0)
    paste_fit(img, A(fig), (x + 30, y + 112, x + col_w - 30, y + 660))
    yy = y + 692
    for s in specs:
        d.rectangle([x + 40, yy + 12, x + 54, yy + 26], fill=TEAL if i == 0 else SAND)
        fit_text(d, (x + 74, yy, x + col_w - 30, yy + 76), s, SANS_R, 30, 22,
                 BODY, 0, line_gap=1.2)
        yy += 82
y += box_h + 72

# ═════════════════════ 7. results ═════════════════════════════════════════
y = section_header(y, "Evaluation", "Results that hold up on unseen data")
stats = [("80%", "ASL top-1, test split"), ("62.4%", "ASL on a foreign corpus (SignASL)"),
         ("99.4%", "ArSL top-1, test split"), ("88%", "ArSL on unseen signers")]
cw = (BW - 2 * M - 3 * 40) // 4
for i, (v, lbl) in enumerate(stats):
    x = M + i * (cw + 40)
    d.rounded_rectangle([x, y, x + cw, y + 224], radius=20, fill=CARD,
                        outline=HAIR, width=3)
    fit_text(d, (x + 16, y + 24, x + cw - 16, y + 134), v, SERIF_B, 92, 60,
             TEAL, 0, align="center")
    fit_text(d, (x + 22, y + 142, x + cw - 22, y + 212), lbl, SANS_R, 29, 21,
             MUTED, 0, align="center", line_gap=1.2)
y += 224 + 44
half = (BW - 2 * M - 60) // 2
r1 = paste_fit(img, A("fig5_1_recognition_accuracy"), (M, y, M + half, y + 545))
r2 = paste_fit(img, A("fig5_4_chrf"), (M + half + 60, y, BW - M, y + 545))
caption((M + half // 2, max(r1[3], r2[3]) + 14), "Recognition: in-distribution vs. generalization")
caption((M + half + 60 + half // 2, max(r1[3], r2[3]) + 14),
        "Translation quality (chrF): the LLM lifts 0.54 → 0.91 (EN), 0.34 → 0.59 (AR)")
y = max(r1[3], r2[3]) + 14 + 40 + 56

# ═════════════════════ 8. the platform ════════════════════════════════════
y = section_header(y, "The platform", "Eight modules, two languages, one web app")
row_h = 470
dash_w = int(row_h * 2.24)
rp = paste_fit(img, A("Figure4_19"), (M, y, M + dash_w, y + row_h), border=True)
rx = rp[2] + 56
rm = paste_fit(img, A("Figure4_24"), (rx, y, rx + int(row_h * 0.53), y + row_h), border=True)
bx = rm[2] + 64
d.text((bx, y + 10), "MODULES", font=F(SANS_B, 28), fill=SAND)
yy = y + 62
for s in ["HandScript · VoiceBridge · SignType",
          "TalkSide · SignLine · Dictionary",
          "Practice · Analytics — 8 in one PWA",
          "Real-time on CPU: 43 ms (ASL) ·\n4.2 ms (ArSL) per inference",
          "Privacy by design: only landmarks\never leave the browser"]:
    d.rectangle([bx + 2, yy + 10, bx + 16, yy + 24], fill=TEAL)
    for j, ln in enumerate(s.split("\n")):
        lf = F(SANS_R, 29)
        d.text((bx + 36, yy + j * 40), ln, font=lf, fill=BODY)
    yy += 46 + 40 * (s.count("\n"))
caption((M + dash_w // 2, y + row_h + 14), "The shared real-time translation dashboard (sign → text, live)")
caption(((rx + rm[2]) // 2, y + row_h + 14), "Mobile PWA — Arabic RTL")
y = y + row_h + 14 + 40

# ═════════════════════ 9. conclusion strip + footer ═══════════════════════
strip_t = BH - 460
print(f"content ends at y={y}, conclusion strip starts at {strip_t} "
      f"(clearance {strip_t - y}px)")
assert y <= strip_t - 10, "content overruns the conclusion strip"
vgrad((0, strip_t, BW, BH - 120), INK, "#0b1830")
logo_at(A("logo_white_hi"), 250, strip_t + 170, 190)
fit_text(d, (440, strip_t + 60, BW - M, strip_t + 200),
         "“An accessible, bidirectional, bilingual sign-language translator is "
         "achievable today — in a browser.”",
         SERIF_B, 54, 36, WHITE, 0, line_gap=1.2)
fit_text(d, (440, strip_t + 215, BW - M, strip_t + 300),
         "Isolated-sign recognition · gloss-mediated LLM translation · semantic sign synthesis · "
         "live meetings — CPU-only, built on public datasets.",
         SANS_R, 32, 24, "#b9c6d8", 0, line_gap=1.25)
d.rectangle([0, BH - 120, BW, BH], fill=WHITE)
d.rectangle([0, BH - 120, BW, BH - 114], fill=SAND)
f = F(SANS_R, 28)
foot = ("Together — Graduation Project · Future University in Egypt · "
        "Faculty of Engineering and Technology · July 2026")
d.text(((BW - text_w(f, foot)) / 2, BH - 84), foot, font=f, fill=MUTED)

img.save(os.path.join(HERE, "Together_Banner.png"))
print("banner PNG saved", img.size)
import deck_engine
print("fit warnings:", deck_engine.WARNINGS or "none")

# ─── package as one-slide PPTX at the exact reference size ───
from pptx import Presentation
from pptx.util import Inches
prs = Presentation()
prs.slide_width = Inches(16.5)
prs.slide_height = Inches(41)
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.shapes.add_picture(os.path.join(HERE, "Together_Banner.png"), 0, 0,
                         width=prs.slide_width, height=prs.slide_height)
prs.save(os.path.join(HERE, "Together_Banner.pptx"))
print("banner PPTX saved (16.5 x 41 in)")
