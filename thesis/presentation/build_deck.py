# -*- coding: utf-8 -*-
"""Together — graduation-defense deck: 70 slides, 4 speakers, ~30 minutes + live demo.
Run:  python3 build_deck.py   → slides/, Together_Defense.pptx, Together_Defense_preview.pdf
"""
import glob
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from deck_engine import render_deck, WARNINGS, ASSETS


def A(name):
    hits = glob.glob(os.path.join(ASSETS, name + ".*"))
    if not hits:
        raise FileNotFoundError(name)
    return hits[0]


def EQ(name):
    return os.path.join(ASSETS, "eq", name + ".png")


S = []  # slides
N = []  # notes


def add(meta, note):
    S.append(meta)
    N.append(note)


# ════════════════════════ SPEAKER 1 — Opening, Problem, Literature ═══════════
add({"type": "title"},
    "[Speaker 1 — 40s] Good morning. We are the Together team, and this is our graduation "
    "project: an AI-based sign-language translator for American and Arabic Sign Language. "
    "Over the next thirty minutes the four of us will walk you through the problem, the two "
    "recognition models we trained, the full web platform we built, and the evaluation — and "
    "then we will show you the system running live.")

add({"type": "cards", "kicker": "Agenda", "title": "Four speakers, one system",
     "section": "Opening", "cols": 2, "cards": [
        ("Speaker 1 — The problem & the field",
         "Why sign-language translation matters, what exists today, and the gap Together fills."),
        ("Speaker 2 — Architecture & the ASL model",
         "System design, the browser-to-server pipeline, and the custom 250-class Conv-Transformer recognizer."),
        ("Speaker 3 — ArSL, language & the platform",
         "The Egyptian ArSL recognizer, the LLM translation layer, sign synthesis, and the eight app modules."),
        ("Speaker 4 — Evaluation & live demo",
         "Recognition and translation results, latency, limitations — then Together running live."),
     ]},
    "[Speaker 1 — 30s] A quick roadmap. I will open with the problem and the research "
    "landscape. My colleagues then cover the system architecture and the ASL model, the "
    "Arabic model and the application itself, and finally the full evaluation. We close "
    "with a live demonstration and your questions.")

add({"type": "stat", "kicker": "The problem", "title": "Hearing loss at global scale",
     "section": "Problem", "stat": "1.5 billion+",
     "context": "people live with some degree of hearing loss — and for millions of them, "
                "a signed language is their first and most natural language.",
     "source": "World Health Organization, Deafness and Hearing Loss fact sheet, 2024 [1]"},
    "[Speaker 1 — 40s] The scale of this problem is easy to underestimate. The WHO counts "
    "more than one and a half billion people with some degree of hearing loss, and the "
    "number is growing. For millions of Deaf people the natural language is a signed one — "
    "ASL, or here in Egypt, Egyptian Arabic Sign Language. The barrier they face daily is "
    "not a disability problem; it is a translation problem.")

add({"type": "content", "kicker": "The problem", "title": "A translation gap, not a linguistic deficit",
     "section": "Problem", "layout": "bottom", "bsize": 30, "images": [A("Figure1_2")],
     "bullets": [
        ("Full languages.", "ASL and ArSL have their own grammar, syntax and space — not signed English or Arabic [2]."),
        ("Two directions needed.", "A real conversation requires sign→speech and speech→sign at the same time — interpreters cannot scale to that."),
     ]},
    "[Speaker 1 — 45s] Sign languages are complete natural languages — Stokoe established "
    "this back in 1960. They are not word-for-word codes of spoken language, which is exactly "
    "why translation is hard. And a real conversation needs both directions at once: the Deaf "
    "participant must be understood, and must understand back. That bidirectional loop, shown "
    "here, is the design target of the whole project.")

add({"type": "cards", "kicker": "The problem", "title": "Why current systems fall short",
     "section": "Problem", "cols": 3, "cards": [
        ("1", "One-way streets", "Systems either read sign or produce sign — almost never both in one platform."),
        ("2", "High-resource bias", "Research concentrates on ASL; Arabic Sign Language is critically under-served."),
        ("3", "Special hardware", "Gloves, depth cameras and multi-camera rigs keep these tools in the lab."),
     ]},
    "[Speaker 1 — 35s] Surveying the field, deployed systems stumble on three hurdles. They "
    "are one-directional. They target high-resource languages, leaving ArSL behind. And they "
    "lean on hardware — gloves, depth cameras, calibrated rigs — that no Deaf person carries "
    "around. These three failures define our requirements.")

add({"type": "content", "kicker": "Our answer", "title": "Together: real-time, two-way, bilingual, in the browser",
     "section": "Problem", "layout": "right", "split": 0.38, "images": [A("Figure1_1")],
     "bullets": [
        ("Zero install.", "A webcam and a browser — nothing else."),
        ("Four directions.", "Sign→text, sign→speech, text→sign, speech→sign."),
        ("Two languages.", "ASL (250 signs) and Egyptian ArSL (20), each with its own model."),
        ("Live meetings.", "Signer and speaker converse over WebRTC."),
     ]},
    "[Speaker 1 — 45s] Together is our answer: a real-time, bidirectional, bilingual "
    "translator that runs entirely in a standard browser. MediaPipe extracts a skeleton of "
    "landmarks client-side, our models classify signs on the server, and a language model "
    "turns sign grammar into fluent sentences. The same machinery runs in reverse to animate "
    "an avatar, and a live meeting mode runs both directions concurrently.")

add({"type": "cards", "kicker": "Objectives", "title": "Six objectives, all delivered",
     "section": "Problem", "cols": 3, "cards": [
        ("O1", "Accessible front-end", "Landmark capture in the browser — no special cameras or hardware."),
        ("O2", "Robust recognition", "Train and test ASL + ArSL models that generalize to unseen signers."),
        ("O3", "Bridge the grammar gap", "LLM turns recognized gloss into fluent sentences, with offline fallback."),
        ("O4", "Reverse translation", "Text and speech back into sign via semantic retrieval and an avatar."),
        ("O5", "Unified experience", "One bilingual web app, including a live two-person meeting mode."),
        ("O6", "Measure the impact", "Quantify recognition accuracy and the LLM's translation gain."),
     ]},
    "[Speaker 1 — 35s] We set six objectives at the start of the year, and every one of them "
    "is delivered in the system you will see today. Keep objective six in mind — measuring "
    "how much the language model actually improves translation — because it shapes the whole "
    "evaluation chapter at the end.")

add({"type": "divider", "num": "01", "title": "The Research Landscape",
     "sub": "Datasets, recognition families, and five commercial systems — and the gap none of them fill."},
    "[Speaker 1 — 10s] Let me place Together in the research landscape first.")

add({"type": "table", "kicker": "Literature", "title": "The data landscape: rich for ASL, scarce for ArSL",
     "section": "Literature", "fsize": 25,
     "headers": ["Dataset", "Language", "Level", "Size", "Signers"],
     "widths": [0.24, 0.18, 0.22, 0.22, 0.14],
     "rows": [
        ["WLASL [6]", "ASL", "Isolated word", "21k videos · 2,000 glosses", "100+"],
        ["MS-ASL [7]", "ASL", "Isolated word", "25k videos · 1,000 signs", "222"],
        ["AUTSL [8]", "Turkish SL", "Isolated word", "38k samples · 226 signs", "43"],
        ["How2Sign [12]", "ASL", "Continuous", "80+ hours", "11"],
        ["PHOENIX14T [10]", "German SL", "Continuous", "8,257 sentences", "9"],
        [("KArSL [13]",), ("ArSL",), "Isolated word", "502 signs · 75k samples", ("3",)],
        [("ArSL2018 [14]",), ("ArSL",), "Static alphabet", "54k images · 32 classes", "40"],
     ]},
    "[Speaker 1 — 45s] The data tells the story. English ASL has tens of thousands of videos "
    "across hundreds of signers. Arabic has KArSL — large but recorded from only three "
    "signers — and alphabet-level image datasets. There is no large word-level Egyptian "
    "corpus, and no continuous ArSL corpus at all. Any honest Arabic system must be designed "
    "for data scarcity, and ours is.")

add({"type": "cards", "kicker": "Literature", "title": "Recognition: from pixels to poses",
     "section": "Literature", "cols": 2, "cards": [
        ("Appearance-based", "2D/3D CNNs (I3D) on raw video — strong but compute-hungry; ≈32% top-1 on 2,000-sign WLASL [6]."),
        ("Pose-based", "Models on extracted skeleton keypoints match appearance models at a fraction of the cost."),
        ("Signer-independent proof", "Skeleton-aware GCNs won the AUTSL challenge at ≈98% [9] — keypoints generalize across people."),
        ("Our choice", "MediaPipe landmarks: compact, private, background-invariant — and they run in a browser."),
     ]},
    "[Speaker 1 — 45s] Recognition research splits into appearance models on raw video and "
    "pose models on extracted keypoints. On large vocabularies raw-video models are expensive "
    "and fragile, while keypoint models proved they generalize across signers — the AUTSL "
    "challenge was won by a skeleton model at ninety-eight percent. That is why Together is "
    "landmark-based end to end: cheaper, more private, and it runs in a browser.")

add({"type": "cards", "kicker": "Literature", "title": "Translation: gloss as the bridge",
     "section": "Literature", "cols": 2, "cards": [
        ("Neural SLT", "Camgöz et al. formalized sign→sentence translation, with gloss as the intermediate representation [10]."),
        ("Sign Language Transformers", "Joint recognition + translation roughly doubled quality on PHOENIX14T [11]."),
        ("The catch", "End-to-end models need large aligned corpora of continuous signing — which do not exist for ArSL."),
        ("Our route", "Isolated signs → gloss → a general-purpose LLM supplies the grammar. No parallel corpus required."),
     ]},
    "[Speaker 1 — 45s] For full translation, the milestone work of Camgöz and colleagues used "
    "gloss — the written shorthand of sign language — as a bridge between video and sentences. "
    "Their transformers are powerful but need thousands of aligned sentences, a resource "
    "Arabic simply does not have. Our route keeps the gloss bridge but replaces the trained "
    "decoder with a general-purpose LLM — no parallel corpus required. The meeting pipeline "
    "shown here is where that choice pays off.")

add({"type": "table", "kicker": "Literature", "title": "Five commercial systems — all one-directional",
     "section": "Literature", "fsize": 24,
     "headers": ["System", "Direction", "Languages", "Hardware", "Key limitation"],
     "widths": [0.16, 0.2, 0.16, 0.22, 0.26],
     "rows": [
        ["SignAll [18]", "Sign → spoken", "ASL", "Multi-camera + gloves", "Fixed rig; not portable"],
        ["Hand Talk [19]", "Spoken → sign", "ASL, Libras", "None (mobile)", "Cannot read signing"],
        ["SLAIT [20]", "Sign → spoken", "ASL", "Webcam", "Small vocabulary; pivoted to lessons"],
        ["KinTrans [21]", "Sign → spoken", "ASL, ArSL", "3-D depth camera", "Fixed installation"],
        ["Signapse [22]", "Spoken → sign", "BSL, ASL", "None (pre-rendered)", "Domain-limited output"],
     ]},
    "[Speaker 1 — 45s] Five commercial systems frame the market. Notice the direction column: "
    "every single one is one-directional. SignAll needs a camera rig and gloves. KinTrans "
    "supports Arabic but requires a depth camera bolted to a counter. SLAIT is closest to us "
    "— browser and webcam — but reads sign only, with a small vocabulary. Nobody closes the "
    "conversational loop.")

add({"type": "cards", "kicker": "Literature", "title": "The gap Together fills",
     "section": "Literature", "cols": 3, "cards": [
        ("Bidirectional", "Both directions in one platform — plus a live meeting mode running them concurrently."),
        ("Bilingual", "ASL and Egyptian ArSL as first-class languages, each with its own trained model."),
        ("Hardware-free", "Webcam + browser only, built on public datasets — fully reproducible."),
     ]},
    "[Speaker 1 — 30s] So the gap is precise: no existing system is bidirectional, bilingual "
    "and hardware-free at the same time. That triple is exactly what Together delivers, and "
    "it is reproducible from public datasets. I will hand over to my colleague to show you "
    "how the system is built.")

# ════════════════════════ SPEAKER 2 — Architecture + ASL model ══════════════
add({"type": "divider", "num": "02", "title": "System Architecture",
     "sub": "A browser client, a FastAPI + Socket.IO server, two recognition models, "
            "and a provider chain that degrades gracefully."},
    "[Speaker 2 — 10s] Thank you. Let me open the hood.")

add({"type": "content", "kicker": "Architecture", "title": "One architecture, four translation paths",
     "section": "Architecture", "layout": "full", "images": [A("F1_architecture")]},
    "[Speaker 2 — 60s] This is the whole system on one slide. In the browser, MediaPipe "
    "Holistic extracts 543 pose, face and hand landmarks per frame — raw video never leaves "
    "the device, which is our privacy stance. Landmarks stream over Socket.IO to a FastAPI "
    "server hosting two recognizers: a TFLite model for ASL and a PyTorch model for Arabic. "
    "A gloss engine and provider chain handle language, PostgreSQL with pgvector handles "
    "semantic sign lookup, and everything cloud has a local fallback — Gemini falls back to "
    "Ollama, cloud TTS to local synthesis.")

add({"type": "content", "kicker": "Architecture", "title": "The two pipelines, end to end",
     "section": "Architecture", "layout": "full", "images": [A("Figure1_2")]},
    "[Speaker 2 — 40s] Reading the top row left to right: camera frames become landmarks, "
    "landmarks become isolated sign predictions, a voting buffer stabilizes them into gloss, "
    "and the LLM composes a sentence which we can also speak aloud. The bottom row is the "
    "mirror: text or transcribed speech is glossed, each token is matched to a stored sign "
    "by semantic search, and the avatar plays the landmark sequences back.")

add({"type": "content", "kicker": "Architecture", "title": "Data flow through six processes",
     "section": "Architecture", "layout": "full", "images": [A("F3_dfd1")]},
    "[Speaker 2 — 35s] The level-one data-flow diagram decomposes that into six processes — "
    "capture, recognition, sentence formation, glossing, sign lookup, and authentication — "
    "with three stores: the sign database with its embeddings, the landmark files on disk, "
    "and the user store. Every arrow here exists as a concrete API route or socket event in "
    "the codebase.")

add({"type": "content", "kicker": "Architecture", "title": "Deployment: CPU-only, container-first",
     "section": "Architecture", "layout": "right", "split": 0.4, "images": [A("F7_deployment")],
     "bullets": [
        ("Client.", "Browser + WASM MediaPipe; installable as a PWA."),
        ("Server.", "FastAPI + Socket.IO on uvicorn, Dockerized."),
        ("Data.", "PostgreSQL 16 + pgvector; landmark store on disk."),
        ("AI providers.", "Gemini primary; Ollama, pyttsx3, Whisper offline."),
        ("~No GPU anywhere.", "Both models are quantized/optimized for CPU."),
     ]},
    "[Speaker 2 — 40s] Deployment is deliberately boring: one container for the web service, "
    "Postgres with the vector extension, and an optional local Ollama. There is no GPU in "
    "this diagram on purpose — both recognizers are optimized to run on commodity CPU, which "
    "is what makes the system deployable on student-budget infrastructure.")

add({"type": "content", "kicker": "Architecture", "title": "Live meetings: translation in both directions at once",
     "section": "Architecture", "layout": "full", "images": [A("Figure1_3")]},
    "[Speaker 2 — 40s] The meeting mode is where everything composes. Two participants "
    "exchange audio and video peer-to-peer over WebRTC — media never transits our server; "
    "Socket.IO only does signalling. The signer's landmarks run through recognition and "
    "reach the speaker as live captions and speech; the speaker's voice is transcribed, "
    "glossed and rendered as an avatar for the signer. Two pipelines, one per participant, "
    "running concurrently.")

add({"type": "content", "kicker": "Architecture", "title": "Streaming recognition as a state machine",
     "section": "Architecture", "layout": "right", "split": 0.42, "images": [A("F13_state")],
     "bullets": [
        ("Rolling buffer.", "Up to 60 frames of landmarks, streamed one frame per tick."),
        ("Vote + cooldown.", "Majority vote over recent predictions; cooldown prevents repeats."),
        ("Hand-gap reset.", "A pause flushes the buffer — the natural sign boundary."),
        ("~Result.", "Noisy per-frame outputs become stable sign tokens."),
     ]},
    "[Speaker 2 — 40s] Real-time recognition is a small state machine. Frames fill a rolling "
    "buffer; inference runs in a worker thread; a prediction only survives if it wins a "
    "majority vote and clears a confidence gate, then a cooldown stops it from repeating. "
    "When the hands drop, the buffer flushes — that pause is the natural boundary between "
    "signs. This is what makes the demo feel stable rather than jittery.")

add({"type": "divider", "num": "03", "title": "The ASL Recognizer",
     "sub": "250 signs · custom Conv1D + Transformer hybrid · trained on Google ISLR · exported to TFLite."},
    "[Speaker 2 — 10s] Now the first of our two models — the large one.")

add({"type": "content", "kicker": "ASL model", "title": "Google ISLR: the training corpus",
     "section": "ASL model", "layout": "right", "split": 0.52, "images": [A("Figure4_1")],
     "bullets": [
        ("Scale.", "≈100,000 landmark sequences · 250 signs · 21 Deaf signers [23]."),
        ("Our split.", "The full corpus, split 70 / 15 / 15 for train, validation, test."),
        ("Landmark-native.", "The dataset ships MediaPipe landmarks — no video processing."),
        ("~Deployed shape.", "TFLite accepts raw (60, 543, 3) and preprocesses inside the graph."),
     ]},
    "[Speaker 2 — 40s] The ASL model trains on Google's Isolated Sign Language Recognition "
    "corpus — about a hundred thousand sequences signed by twenty-one Deaf adults. We train "
    "on the full corpus under a seventy-fifteen-fifteen split. Note the deployment "
    "detail on the right: the exported TFLite graph takes raw landmarks and performs all "
    "preprocessing internally, so the browser sends coordinates and nothing else.")

add({"type": "content", "kicker": "ASL model", "title": "Preprocessing: 543 points → 708 features",
     "section": "ASL model", "layout": "bottom", "bsize": 30, "images": [A("Figure4_2")],
     "bullets": [
        ("Select & center.", "118 informative landmarks (lips, hands, upper body); drop z; center on the nose."),
        ("Normalize & differentiate.", "Scale by per-clip σ; add velocity Δ¹ and acceleration Δ² → 118 × 6 = 708 features."),
     ]},
    "[Speaker 2 — 45s] Preprocessing distills the 543 raw points down to the 118 that carry "
    "signing information — lips, both hands, upper body. We drop depth, center on the nose, "
    "and normalize by the clip's standard deviation so distance from the camera cancels out. "
    "Then we append first- and second-order motion — velocity and acceleration — giving 708 "
    "features per frame. Motion features matter because a sign is a trajectory, not a pose.")

add({"type": "content", "kicker": "ASL model", "title": "Macro architecture: Conv1D × 3 + Transformer, twice",
     "section": "ASL model", "layout": "right", "split": 0.42, "images": [A("Figure4_3")],
     "bullets": [
        ("Stem.", "Linear projection of 708 features to a 192-d representation."),
        ("Body.", "Two repetitions of three Conv1D blocks followed by a Transformer block."),
        ("Head.", "Global average pooling → 250-way softmax."),
        ("~Design.", "A custom hybrid built for this task: convolution for local shape, attention for long range."),
     ]},
    "[Speaker 2 — 40s] The architecture is our own convolution-plus-transformer hybrid: "
    "convolutions capture local hand-shape transitions, attention captures the "
    "long-range structure of a sign, and we alternate them — three convolution blocks then a "
    "transformer block, repeated twice — before pooling into a 250-way classifier.")

add({"type": "content", "kicker": "ASL model", "title": "Inside a Conv1D block — causal by construction",
     "section": "ASL model", "layout": "bottom", "bsize": 29,
     "images": [A("Figure4_4"), A("Figure4_7")],
     "bullets": [
        ("Structure.", "Point-wise expansion → causal depthwise conv (k=17) → BatchNorm → channel attention → projection, wrapped in a residual."),
        ("~Causal padding.", "No future frames leak into the past — the same network is safe for live streaming."),
     ]},
    "[Speaker 2 — 45s] Each convolution block is an inverted bottleneck: a point-wise "
    "expansion, a depthwise convolution with a wide kernel of seventeen frames, channel "
    "attention, and a residual connection. The detail worth defending is causal padding — "
    "shown lower right — which guarantees the model never looks at future frames. That is "
    "what lets the same network run on a live stream without behavioral drift.")

add({"type": "content", "kicker": "ASL model", "title": "Efficient Channel Attention: focus for almost nothing",
     "section": "ASL model", "layout": "bottom", "bsize": 30, "images": [A("Figure4_5")],
     "bullets": [
        ("Squeeze.", "Global-average-pool each of the 384 channels to a single value."),
        ("Excite.", "A 1-D convolution across channels yields sigmoid weights that rescale the features."),
     ]},
    "[Speaker 2 — 30s] Efficient Channel Attention is a five-line idea: pool each channel to "
    "one number, run a tiny one-dimensional convolution across channels, and use the sigmoid "
    "output to reweight them. It costs a few hundred parameters and consistently sharpens "
    "which feature channels the network trusts at each moment.")

add({"type": "content", "kicker": "ASL model", "title": "Transformer block + a three-tower ensemble",
     "section": "ASL model", "layout": "right", "split": 0.42, "images": [A("Figure4_6")],
     "bullets": [
        ("Attention.", "Multi-head self-attention: H = 4 heads, d_k = 48, pre-norm residuals."),
        ("FFN.", "Two-layer MLP with swish activation."),
        ("~Ensemble.", "Three independently-trained towers averaged at the logit level in the deployed TFLite."),
     ]},
    "[Speaker 2 — 35s] The transformer block is the standard pre-norm design — four heads of "
    "dimension forty-eight with a feed-forward network. One deployment choice: the shipped "
    "TFLite file is actually a three-tower ensemble, three independently trained copies "
    "averaged at the logit level. It buys robustness at inference time for zero extra "
    "engineering in the client.")

add({"type": "content", "kicker": "ASL model", "title": "Augmentation: train for messy reality",
     "section": "ASL model", "layout": "bottom", "bsize": 30, "images": [A("Figure4_8")],
     "bullets": [
        ("Temporal.", "Resample 0.5–1.5× for signing speed; mask 20–40% of frames for dropouts."),
        ("Spatial.", "Horizontal flip for left-handed signers; affine transforms to ±30°; random cutout."),
     ]},
    "[Speaker 2 — 35s] Augmentation is where generalization is won. Temporal resampling "
    "simulates fast and slow signers; frame masking simulates tracking dropouts; horizontal "
    "flips cover left-handed signers; affine jitter and cutout simulate camera angle and "
    "occlusion. Every effect we expect from a real webcam is manufactured during training.")

add({"type": "content", "kicker": "ASL model", "title": "Training: heavy regularization for 250 classes",
     "section": "ASL model", "layout": "bottom", "bsize": 30, "images": [A("Figure4_9")],
     "bullets": [
        ("Optimizer.", "RAdam + Lookahead, cosine decay from 4e-3, 400 epochs, label smoothing ε = 0.1."),
        ("Regularizers.", "Drop-Path 0.2 · late dropout 0.8 · Adversarial Weight Perturbation λ = 0.2."),
     ]},
    "[Speaker 2 — 40s] Training runs four hundred epochs with RAdam wrapped in Lookahead "
    "under a cosine schedule, and cross-entropy with label smoothing. Against overfitting on "
    "250 classes we stack three regularizers: stochastic depth, a very aggressive dropout on "
    "the final layer, and adversarial weight perturbation, which optimizes the loss under "
    "small worst-case weight shifts.")

add({"type": "content", "kicker": "ASL model", "title": "Learning curves: validation reaches 0.80",
     "section": "ASL model", "layout": "bottom",
     "images": [A("fig4_10a_asl_accuracy"), A("fig4_10b_asl_loss")]},
    "[Speaker 2 — 40s] The curves tell a healthy story: validation accuracy climbs to about "
    "eighty percent and plateaus with no divergence between the curves. If validation above "
    "training looks odd — that is expected here, because the heavy augmentation and "
    "regularization are active only on the training pass. Speaker three will take it from "
    "here with the Arabic model.")

# ════════════════════════ SPEAKER 3 — ArSL + Language + App ═════════════════
add({"type": "divider", "num": "04", "title": "The ArSL Recognizer",
     "sub": "20 Egyptian signs · compact CNN + bidirectional GRU · built for a low-resource language."},
    "[Speaker 3 — 10s] Thank you. The Arabic model had the opposite problem: not too many classes — too little data.")

add({"type": "cards", "kicker": "ArSL model", "title": "Balaha ArSL-20: small, but many signers",
     "section": "ArSL model", "cols": 3, "cards": [
        ("8,437 samples", "Phone-camera clips of 20 Egyptian signs — realistic, noisy capture conditions [24]."),
        ("72 signers", "The real asset: enough signer diversity to hold entire people out of training."),
        ("70 / 15 / 15", "Stratified split for training — plus a separate signer-independent protocol in the evaluation."),
     ]},
    "[Speaker 3 — 40s] Our Arabic corpus is the Balaha twenty-word dataset: eight and a half "
    "thousand phone-camera clips. Twenty signs is small — but seventy-two different signers "
    "is the valuable part, because it lets us later hold entire signers out of training and "
    "measure honest generalization. The architecture must respect the data budget: a "
    "transformer here would simply memorize.")

add({"type": "content", "kicker": "ArSL model", "title": "Preprocessing: invariance in four steps",
     "section": "ArSL model", "layout": "bottom", "bsize": 30, "images": [A("Figure4_11")],
     "bullets": [
        ("59 landmarks.", "17 upper-body pose points + 21 per hand; z zeroed — phone depth is noise."),
        ("Normalize + resample.", "Center on the shoulder midpoint, scale by shoulder width, resample to 30 frames → (30, 177)."),
     ]},
    "[Speaker 3 — 35s] Preprocessing keeps fifty-nine landmarks — upper body and both hands — "
    "and zeroes the depth axis, which is unreliable across phone cameras. Centering on the "
    "shoulder midpoint and scaling by shoulder width buys position and size invariance, and "
    "every clip is resampled to exactly thirty frames, so the network always sees a thirty-by-"
    "one-seventy-seven tensor.")

add({"type": "content", "kicker": "ArSL model", "title": "A compact CNN-GRU, sized to the data",
     "section": "ArSL model", "layout": "right", "split": 0.4, "images": [A("Figure4_12")],
     "bullets": [
        ("Convolutions.", "Two 1-D blocks (177→128→128, kernel 3) with BatchNorm + ReLU."),
        ("Recurrence.", "Two-layer bidirectional GRU, 64 hidden units, dropout 0.3."),
        ("Head.", "128 → 64 → 20 with dropout 0.5 and softmax."),
        ("~At inference.", "Soft-max averaged over 30/45/60-frame windows; accept above τ = 0.45."),
     ]},
    "[Speaker 3 — 40s] The model is a deliberately compact hybrid: two small convolution "
    "blocks learn local motion patterns, and a two-layer bidirectional GRU models the "
    "temporal order. At inference we run three temporal windows — thirty, forty-five and "
    "sixty frames — and average their probabilities before a confidence gate, which smooths "
    "out any single bad window.")

add({"type": "content", "kicker": "ArSL model", "title": "Why bidirectional recurrence",
     "section": "ArSL model", "layout": "bottom", "images": [A("Figure4_13")],
     "bullets": [
        ("Full context.", "Forward and backward passes are concatenated — every prediction sees the whole sign."),
     ]},
    "[Speaker 3 — 30s] The GRU reads each sequence in both directions and concatenates the "
    "two hidden states, so the decision at the end of a sign is informed by its beginning "
    "and vice versa. For isolated signs — where we always have the full clip in the buffer — "
    "bidirectionality is free accuracy.")

add({"type": "content", "kicker": "ArSL model", "title": "Augmentation and training, tuned for small data",
     "section": "ArSL model", "layout": "bottom", "bsize": 29,
     "images": [A("Figure4_14"), A("Figure4_15")],
     "bullets": [
        ("Online augmentation.", "Mirroring 50% · inactive-hand masking 70% · affine 0.92–1.08× + Gaussian noise · temporal jitter."),
        ("Training.", "Adam (wd 1e-4), lr 1e-3 with ReduceLROnPlateau, batch 64, ≤100 epochs, early stopping (patience 12)."),
     ]},
    "[Speaker 3 — 35s] Augmentation mirrors half the samples, masks the inactive hand — "
    "because one-handed signs should not depend on where the idle hand rests — and adds "
    "affine scaling with a little noise plus temporal jitter. Training is plain Adam with a "
    "plateau scheduler and early stopping; the model converges in well under a hundred epochs.")

add({"type": "content", "kicker": "ArSL model", "title": "Learning curves: 99.41% on the test split",
     "section": "ArSL model", "layout": "bottom",
     "images": [A("fig4_16a_arsl_accuracy"), A("fig4_16b_arsl_loss")]},
    "[Speaker 3 — 30s] On its own test split the model converges fast and lands at ninety-"
    "nine point four percent. Hold your skepticism — a stratified split shares signers across "
    "partitions, so speaker four will show the honest signer-independent number later. Even "
    "there, this small model holds up remarkably well.")

add({"type": "divider", "num": "05", "title": "From Signs to Sentences",
     "sub": "Gloss is not grammar. A language-model layer bridges Topic–Comment sign order and fluent sentences — both ways."},
    "[Speaker 3 — 10s] Recognition gives us words. Language is the next problem.")

add({"type": "cards", "kicker": "Language layer", "title": "The grammar gap",
     "section": "Language", "cols": 3, "cards": [
        ("Topic–Comment order", "ASL fronts the topic: “STORE I GO” — not “I am going to the store.”"),
        ("Dropped words", "Articles and copulas simply do not exist in sign gloss."),
        ("Facial grammar", "Questions and negation live on the face — non-manual markers [16], [17]."),
     ]},
    "[Speaker 3 — 40s] Sign languages order information Topic-first: the gloss for “I am "
    "going to the store” is literally STORE — I — GO, with no articles and no copulas. And a "
    "third of the grammar is not on the hands at all: raised eyebrows mark yes-no questions, "
    "furrowed brows mark wh-questions. Concatenating recognized words can never produce a "
    "fluent sentence — something must supply the grammar.")

add({"type": "cards", "kicker": "Language layer", "title": "An LLM supplies the grammar — with a safety net",
     "section": "Language", "cols": 3, "cards": [
        ("1", "Gemini 3.5 Flash", "Few-shot prompts enforce Topic–Comment ↔ SVO conversion in both languages [27]."),
        ("2", "Bounded cache", "256 recent gloss→sentence results — repeated phrases cost ~0 ms."),
        ("3", "Offline fallback", "No cloud? Ollama serves a local LLaMA [28]; worst case, raw gloss is shown."),
     ]},
    "[Speaker 3 — 40s] We hand the gloss to a large language model with few-shot examples "
    "that teach the reordering. Three engineering guarantees make this production-grade: a "
    "bounded cache so conversation phrases stop costing round-trips, an automatic fallback "
    "to a local model when the cloud is unreachable, and a final fallback to raw gloss — the "
    "feature degrades, it never dies.")

add({"type": "cards", "kicker": "Language layer", "title": "The reverse path — and facial grammar as data",
     "section": "Language", "cols": 3, "cards": [
        ("Sentence → gloss", "“What is your name?” becomes YOUR NAME WHAT — same few-shot layer, reversed."),
        ("NMM annotation", "Sentence type is detected and mapped to markers: eyebrows furrowed, head forward."),
        ("Render-ready grammar", "The avatar layer receives facial grammar it can act on — not just a word list."),
     ]},
    "[Speaker 3 — 35s] The same layer runs in reverse for text-to-sign: sentences become "
    "gloss tokens, and we annotate the non-manual markers — the system knows a wh-question "
    "should carry furrowed brows across its span. Those annotations ship with the animation "
    "data, so facial grammar is represented explicitly rather than lost.")

add({"type": "cards", "kicker": "Sign synthesis", "title": "Finding the right sign: semantic retrieval",
     "section": "Synthesis", "cols": 3, "cards": [
        ("Embed", "Every stored sign carries a Sentence-BERT embedding (all-MiniLM-L6-v2, 384-d) [26]."),
        ("Search", "pgvector HNSW cosine index inside PostgreSQL — a paraphrase still finds the sign."),
        ("Gate", "Language-specific distance thresholds: 0.35 for English, 0.55 for Arabic."),
     ]},
    "[Speaker 3 — 35s] To animate a word we search semantically, not by string match. Every "
    "sign in the database has a sentence-embedding; the query is embedded and matched by "
    "cosine distance inside Postgres itself, using an HNSW index. Ask for “home” and the "
    "sign for “house” answers. A distance gate per language keeps bad matches out.")

add({"type": "cards", "kicker": "Sign synthesis", "title": "Out of vocabulary? Spell it — smoothly",
     "section": "Synthesis", "cols": 3, "cards": [
        ("Fingerspell", "Unknown words decompose into per-letter sign clips."),
        ("Bridge", "5 interpolated frames inserted at every clip boundary."),
        ("Smooth", "A Savitzky–Golay filter (window 7) erases the stitch jerk, preserving each letter."),
     ]},
    "[Speaker 3 — 30s] When a word has no dedicated sign, we synthesize it: fingerspell "
    "letter by letter, insert short interpolated bridges between clips, and run a Savitzky-"
    "Golay filter over each boundary so the concatenation reads as one continuous motion "
    "instead of a slideshow of letters.")

add({"type": "content", "kicker": "Sign synthesis", "title": "The avatar: landmarks back to motion",
     "section": "Synthesis", "layout": "right", "split": 0.5, "images": [A("Figure4_20"), ],
     "border": True,
     "bullets": [
        ("Canvas renderer.", "Retrieved landmark sequences drive a skeleton avatar in the browser."),
        ("Reference videos.", "Each dictionary entry pairs the avatar with the original human signer."),
        ("~Same store, both ways.", "The landmarks we recognize from are the landmarks we animate with."),
     ]},
    "[Speaker 3 — 35s] The avatar is a canvas renderer that replays those landmark sequences "
    "as a moving skeleton — here inside the dictionary, side by side with the original human "
    "recording so learners can compare. One elegant property: the exact representation we "
    "recognize from is the representation we animate with. Now let me show you the product "
    "this all lives in.")

add({"type": "divider", "num": "06", "title": "The Together Platform",
     "sub": "Eight modules, two languages, one design system — installable as a PWA on desktop and mobile."},
    "[Speaker 3 — 10s] Everything so far ships inside one bilingual web application.")

add({"type": "cards", "kicker": "The platform", "title": "Eight modules",
     "section": "Platform", "cols": 4, "cards": [
        ("HandScript", "Sign → text"),
        ("VoiceBridge", "Sign → speech"),
        ("SignType", "Text → sign avatar"),
        ("TalkSide", "Speech → sign avatar"),
        ("SignLine", "Live two-party meeting"),
        ("Dictionary", "Browse & search all signs"),
        ("Practice", "Avatar-prompted drills"),
        ("Analytics", "On-device session stats"),
     ]},
    "[Speaker 3 — 30s] Five translation modules cover every direction plus the live meeting, "
    "and three companion modules — Dictionary, Practice and Analytics — turn the platform "
    "into a learning tool as well as a translator. All eight share one real-time dashboard "
    "and one design system.")

add({"type": "content", "kicker": "The platform", "title": "The shared real-time dashboard",
     "section": "Platform", "layout": "full", "images": [A("Figure4_19")], "border": True},
    "[Speaker 3 — 35s] This is HandScript translating live: camera panel with start-stop, "
    "recognized signs streaming into the session log, a confidence ring updating per "
    "inference, and the transcription panel where accumulated gloss becomes a sentence — "
    "automatically on a pause, or on demand. The status cards along the top track today's "
    "signs, average confidence and session time.")

add({"type": "content", "kicker": "The platform", "title": "Dictionary: avatar + human, side by side",
     "section": "Platform", "layout": "full", "images": [A("Figure4_20")], "border": True},
    "[Speaker 3 — 25s] The Dictionary exposes the whole vocabulary with category filters and "
    "semantic search — a paraphrase still lands on the right entry. Every sign renders as "
    "the avatar and, where available, the original signer video, which learners told us is "
    "the combination they actually want.")

add({"type": "content", "kicker": "The platform", "title": "Analytics: insight without surveillance",
     "section": "Platform", "layout": "full", "images": [A("Figure4_21")], "border": True},
    "[Speaker 3 — 25s] Analytics summarizes detections, unique signs, confidence trends and "
    "signing speed — computed entirely on the device, consistent with our privacy rule that "
    "only landmarks ever leave the browser. Sessions export to CSV or print to PDF.")

add({"type": "content", "kicker": "The platform", "title": "Practice: drilled by the real recognizer",
     "section": "Platform", "layout": "full", "images": [A("Figure4_22")], "border": True},
    "[Speaker 3 — 25s] Practice runs ten-sign rounds: the avatar demonstrates, the learner "
    "attempts, and the very same recognition pipeline that powers translation judges the "
    "attempt with live feedback. You are graded by the production system, not a toy.")

add({"type": "content", "kicker": "The platform", "title": "Fully bilingual, fully responsive",
     "section": "Platform", "layout": "bottom", "border": True,
     "images": [A("Figure4_23"), A("Figure4_24")],
     "bullets": [
        ("True RTL.", "The Arabic interface is a complete right-to-left mirror — every label, chart and control."),
     ]},
    "[Speaker 3 — 25s] The Arabic experience is not a translated skin — it is a true right-"
    "to-left mirror of the entire interface, down to charts and the avatar panel, with no "
    "English leaking through. The same views adapt to mobile, where Together installs as a "
    "progressive web app.")

add({"type": "cards", "kicker": "The platform", "title": "Security & privacy, end to end",
     "section": "Platform", "cols": 3, "cards": [
        ("Argon2id", "Memory-hard password hashing; only hashes are stored."),
        ("Rotating JWT", "Short-lived access tokens; refresh tokens hashed, expiring, rotated on use."),
        ("WhatsApp OTP", "Registration verified by a one-time code — accounts tie to a real number."),
        ("Rate limiting", "Sliding-window limits blunt credential stuffing."),
        ("CORS + headers", "Restrictive origin policy and standard hardening headers."),
        ("Landmarks only", "Raw video never leaves the browser — recognition sees coordinates."),
     ]},
    "[Speaker 3 — 35s] Security is thesis-grade, not an afterthought: Argon2id hashing, "
    "rotating refresh tokens, sliding-window rate limits, and registration verified by a "
    "WhatsApp one-time password. And the deepest privacy property is architectural — the "
    "server never receives video, only landmark coordinates. Speaker four will now put "
    "numbers on all of this.")

# ════════════════════════ SPEAKER 4 — Evaluation, Closing, Demo ═════════════
add({"type": "divider", "num": "07", "title": "Does It Work?",
     "sub": "Recognition, generalization, translation quality against human references, and latency — with baselines."},
    "[Speaker 4 — 10s] Thank you. Let us put every claim under a number.")

add({"type": "cards", "kicker": "Evaluation", "title": "Four families of evidence",
     "section": "Evaluation", "cols": 2, "cards": [
        ("Recognition", "Accuracy, macro precision / recall / F1, confusion analysis — on held-out test sets."),
        ("Generalization", "ASL: cross-dataset on SignASL. ArSL: held-out signers never seen in training."),
        ("Translation quality", "BLEU and chrF versus human references, always against a raw-gloss baseline [32]–[34]."),
        ("System performance", "Per-model inference latency and the effect of INT8 quantization on CPU."),
     ]},
    "[Speaker 4 — 35s] Our methodology rests on four families of evidence. The one I want to "
    "highlight is generalization: models are easy to flatter with a random split, so we "
    "tested the ASL model on an entirely different corpus, and the Arabic model on signers "
    "it had never seen. And translation quality is always measured against the raw-gloss "
    "baseline, so the language model's contribution is isolated, not asserted.")

add({"type": "content", "kicker": "Evaluation", "title": "Recognition: in-distribution vs. the honest number",
     "section": "Evaluation", "layout": "full",
     "images": [A("fig5_1_recognition_accuracy")]},
    "[Speaker 4 — 45s] Here are the four headline numbers. ASL: eighty percent on its own "
    "test distribution, and sixty-two point four when we take the same model to SignASL — a "
    "different corpus, different signers, different recording conditions. That drop is "
    "expected; still recognizing the majority of two hundred and fifty signs on a foreign "
    "corpus is evidence of real sign representations. ArSL: ninety-nine point four in-"
    "distribution, and eighty-eight percent on held-out signers — the honest number, and a "
    "strong one for phone-camera video.")

add({"type": "table", "kicker": "Evaluation", "title": "Per-model metrics on held-out test sets",
     "section": "Evaluation", "fsize": 27,
     "headers": ["Metric", "ASL (250-class)", "ArSL (20-class)"],
     "widths": [0.44, 0.28, 0.28],
     "rows": [
        ["Top-1 accuracy", ("80.00%",), ("99.41%",)],
        ["Top-5 accuracy", "95.00%", "100.00%"],
        ["Macro precision", "0.81", "0.99"],
        ["Macro recall", "0.80", "0.99"],
        ["Macro F1", "0.80", "0.99"],
        ["Cross-entropy loss", "0.90", "0.03"],
     ]},
    "[Speaker 4 — 30s] The full metric tables. Macro precision and recall track accuracy "
    "closely in both models, which tells you performance is balanced across classes rather "
    "than carried by a few easy signs. Top-five accuracy of ninety-five percent for ASL "
    "matters in practice: the right sign is almost always among the top candidates feeding "
    "our voting buffer.")

add({"type": "content", "kicker": "Evaluation", "title": "ArSL: where the 0.6% lives",
     "section": "Evaluation", "layout": "full", "images": [A("Figure5_2")]},
    "[Speaker 4 — 30s] The full twenty-by-twenty confusion matrix for Arabic is almost "
    "perfectly diagonal — the handful of off-diagonal cells are single confusions between "
    "visually adjacent signs. There is no systematic failure mode here; at this vocabulary "
    "size the model is essentially solved in-distribution.")

add({"type": "content", "kicker": "Evaluation", "title": "ASL: which signs fool the model — and why",
     "section": "Evaluation", "layout": "full",
     "images": [A("fig5_3_confusion_analysis")]},
    "[Speaker 4 — 45s] For ASL we ranked misclassifications by the model's confidence in the "
    "wrong sign, from a re-run of the cross-dataset evaluation. Look at what confuses it: "
    "listen and hear. Cut and scissors — in both directions. Chin and thirsty, both "
    "articulated at the chin. These are signs that genuinely look alike; the model's errors "
    "are the errors a human learner makes, which is exactly what a credible recognizer's "
    "failure profile should look like.")

add({"type": "content", "kicker": "Evaluation", "title": "Translation quality: the LLM's contribution, isolated",
     "section": "Evaluation", "layout": "full", "images": [A("fig5_4_chrf")]},
    "[Speaker 4 — 40s] Now objective six: does the language model actually help? Character-"
    "level F-score against human references: for English, raw gloss scores 0.54 and the LLM "
    "lifts it to 0.91. For Arabic, 0.34 to 0.59. Both gains are large and consistent — and "
    "the Arabic numbers are structurally lower not because that system is worse, but because "
    "Arabic morphology spreads correct meaning across many surface forms that overlap "
    "metrics penalize.")

add({"type": "content", "kicker": "Evaluation", "title": "BLEU-n profiles: fluency, not just words",
     "section": "Evaluation", "layout": "full", "images": [A("fig5_5_bleu_profiles")]},
    "[Speaker 4 — 35s] The BLEU profiles make the mechanism visible. Raw gloss holds some "
    "unigram credit — the content words are there — but collapses toward zero at four-grams: "
    "gloss has words, not phrasing. With the LLM the curves stay high across all n-gram "
    "orders: BLEU-4 of 0.81 for English and 0.42 for Arabic, against baselines of 0.04 and "
    "0.03. A twenty-fold improvement in fluent phrasing.")

add({"type": "content", "kicker": "Evaluation", "title": "Precision vs. recall: the LLM supplies grammar",
     "section": "Evaluation", "layout": "full", "images": [A("fig5_6_precision_recall")]},
    "[Speaker 4 — 35s] Decomposing chrF nails down what the LLM adds. Raw gloss already has "
    "decent precision — its words are correct — but poor recall: everything grammatical is "
    "missing. The long arrows are recall: for English from 0.52 to 0.91. The language model "
    "is not polishing word choice; it is supplying the function words and morphology — the "
    "grammar — that gloss by definition lacks.")

add({"type": "cards", "kicker": "Evaluation", "title": "Fast enough to feel live — on CPU",
     "section": "Evaluation", "cols": 4, "cards": [
        ("43.5 ms", "ASL inference, average (58.7 ms max)"),
        ("4.2 ms", "ArSL inference after INT8 — 1.78× faster than fp32"),
        ("~0 ms", "Repeated gloss→sentence via the 256-entry cache"),
        ("0 GPUs", "TFLite XNNPACK + quantized PyTorch, 2–4 GB RAM total"),
     ]},
    "[Speaker 4 — 35s] Latency: the ASL ensemble answers in about forty-three milliseconds "
    "and the quantized Arabic model in four — INT8 quantization bought a 1.78-times speedup "
    "with no measurable accuracy change. The gloss cache removes repeated LLM round-trips "
    "entirely. All of it on CPU: the demo you are about to see runs on hardware like this "
    "podium laptop.")

add({"type": "cards", "kicker": "Honesty", "title": "Limitations we want on the record",
     "section": "Limitations", "cols": 2, "cards": [
        ("Isolated signs, not continuous", "The LLM composes sentences from isolated signs; natural conversational signing is future work."),
        ("ASL split is random", "The 80% is not signer-independent; the 62.4% cross-dataset result is the honest proxy."),
        ("20-word Arabic vocabulary", "A real contribution to a low-resource language — but far from practical coverage."),
        ("20 reference sentences", "BLEU on a small set is high-variance; chrF and human judgement carry more weight."),
     ]},
    "[Speaker 4 — 40s] Four limitations, stated plainly. We recognize isolated signs, not "
    "continuous signing. The ASL in-distribution split shares signers, which is why we lean "
    "on the cross-dataset number. The Arabic vocabulary is twenty words. And our translation "
    "reference set is small, so we weight chrF over BLEU. We would rather you hear these "
    "from us than find them yourselves.")

add({"type": "cards", "kicker": "Honesty", "title": "Each limitation maps to a next step",
     "section": "Limitations", "cols": 2, "cards": [
        ("Isolated → continuous", "Sequence-to-sequence recognition over the same landmark stream."),
        ("Random → signer-independent", "Re-train and re-evaluate ASL under held-out signers, as done for ArSL."),
        ("20 → hundreds of signs", "A documented, consented Egyptian ArSL collection effort."),
        ("Metrics → people", "A user study with Deaf and hard-of-hearing participants on real tasks."),
     ]},
    "[Speaker 4 — 25s] None of these is a dead end — each maps to a concrete next step, and "
    "the landmark-based architecture carries over unchanged to all four. That is, in our "
    "view, the mark of a good foundation.")

add({"type": "cards", "kicker": "Conclusion", "title": "Four contributions",
     "section": "Conclusion", "cols": 2, "cards": [
        ("1", "A bidirectional, bilingual, browser-based translator",
         "Four directions plus live meetings — no gloves, no depth cameras, no installation."),
        ("2", "A working Egyptian ArSL recognizer",
         "88% signer-independent — a concrete contribution to an under-resourced language."),
        ("3", "Gloss-mediated, LLM-assisted translation",
         "Fluent output without the parallel corpora end-to-end models require — chrF 0.54 → 0.91."),
        ("4", "A reproducible engineering blueprint",
         "Public datasets, CPU-only inference, graceful offline degradation."),
     ]},
    "[Speaker 4 — 40s] Our contributions, in order of what we believe matters: the first "
    "system in this space that is bidirectional, bilingual and hardware-free at once; a "
    "working Egyptian sign recognizer validated on held-out signers; the demonstration that "
    "a general-purpose LLM can replace the parallel corpora this field usually requires; and "
    "a blueprint anyone can reproduce from public data on commodity hardware.")

add({"type": "cards", "kicker": "Conclusion", "title": "Where this goes next",
     "section": "Conclusion", "cols": 3, "cards": [
        ("Continuous signing", "From isolated signs to sentence-level recognition."),
        ("Bigger ArSL", "Expand the Egyptian vocabulary; document and consent the corpus."),
        ("Signer-independent ASL", "The same honest protocol we applied to Arabic."),
        ("User study", "Deaf participants, real tasks, comprehension and usability."),
        ("Edge deployment", "Recognition fully in-browser — cut the server round-trip."),
        ("More languages", "The architecture already spans two typologically different ones."),
     ]},
    "[Speaker 4 — 25s] Future work follows directly: continuous signing, a larger consented "
    "Arabic corpus, signer-independent ASL evaluation, a study with the Deaf community, "
    "fully on-device recognition, and new languages — the architecture has already shown it "
    "generalizes across two very different ones.")

add({"type": "closing",
     "title": "“An accessible, bidirectional, bilingual sign-language translator is achievable today.”",
     "sub": "Commodity hardware. Public datasets. A browser. Together is our proof — and our contribution "
            "to technology that serves under-represented Deaf communities.",
     "foot": "Together — FUE Faculty of Engineering and Technology · July 2026"},
    "[Speaker 4 — 25s] Our thesis closes with this sentence, and we stand behind every word "
    "of it: an accessible, bidirectional, bilingual sign-language translator is achievable "
    "today. Not in a lab — in a browser. Let us show you.")

add({"type": "divider", "num": "▶", "title": "Live Demo",
     "sub": "1 · HandScript: live ASL sign → sentence → speech.   2 · SignType: type a sentence → the avatar signs it.   "
            "3 · Arabic dashboard (RTL) → ArSL recognition.   4 · SignLine: two-device live meeting."},
    "[Speaker 4 — 4-5 min] Demo runbook: start on the English dashboard, sign two or three "
    "ASL signs, let the vote buffer commit them, and trigger sentence formation — then play "
    "the spoken output. Switch to SignType and type a sentence for the avatar. Flip the "
    "language toggle to show the full Arabic right-to-left interface and recognize one ArSL "
    "sign. If time allows, open SignLine on the second laptop and hold a two-way exchange. "
    "Fallback if the network misbehaves: the offline chain — Ollama and local TTS — keeps "
    "every step functional.")

add({"type": "closing", "title": "Thank you.",
     "sub": "Questions welcome — on the models, the language layer, the platform, or the evaluation.",
     "foot": "Abdelfattah Moustafa · Michael Abdallah · Ahmed Mohamed Nagib · Ahmed Ashraf Shawareb — Supervisor: Prof. Medhat Awadallah"},
    "[All — Q&A] Thank you for your attention. We are happy to take questions in any order — "
    "each of us will field the section they presented. The backup slide that follows holds "
    "the full hyperparameter tables if the committee wants exact training configurations.")

add({"type": "table", "kicker": "Backup — for Q&A", "title": "Full training configurations",
     "section": "Backup", "fsize": 24,
     "headers": ["Hyperparameter", "ASL (Conv1D + Transformer)", "ArSL (CNN-BiGRU)"],
     "widths": [0.3, 0.37, 0.33],
     "rows": [
        ["Dataset / classes", "Google ISLR — 250 signs [23]", "Balaha ArSL-20 — 20 signs [24]"],
        ["Input tensor", "(384, 708) train · raw (60, 543, 3) deployed", "(30, 177)"],
        ["Optimizer / LR", "RAdam + Lookahead · 5e-4→4e-3 cosine", "Adam (wd 1e-4) · 1e-3 ReduceLROnPlateau"],
        ["Epochs / batch", "400", "≤100, early stop (pat. 12) · batch 64"],
        ["Loss", "CCE + label smoothing ε = 0.1", "Cross-entropy"],
        ["Regularization", "Drop-Path 0.2 · dropout 0.8 · AWP λ = 0.2", "Dropout 0.3 (GRU) / 0.5 (dense)"],
        ["Inference gate", "τ = 0.80 + 15-vote majority", "τ = 0.45 + 30/45/60 window ensemble"],
        ["Export", "TFLite (3-tower ensemble, XNNPACK)", "PyTorch, INT8 dynamic quantization"],
     ]},
    "[Backup — not presented] Complete training configurations for both models, kept for "
    "committee questions. Everything here matches Chapter 4 of the thesis exactly; if asked "
    "about any single value, the relevant thesis section is 4.2.4 for ASL and 4.3.4 for ArSL.")

import re as _re


def _strip_cites(o):
    if isinstance(o, str):
        return _re.sub(r"\s*\[\d+\](\s*[–-]\s*\[\d+\])?(,\s*\[\d+\])*", "", o)
    if isinstance(o, tuple):
        return tuple(_strip_cites(x) for x in o)
    if isinstance(o, list):
        return [_strip_cites(x) for x in o]
    if isinstance(o, dict):
        return {k: (v if k in ("images", "eq") else _strip_cites(v)) for k, v in o.items()}
    return o


S = [_strip_cites(m) for m in S]

assert len(S) == 70, f"slide count = {len(S)}"

# ── render ──────────────────────────────────────────────────────────────────
for i, meta in enumerate(S, 1):
    meta.setdefault("section", meta.get("section", ""))
paths = render_deck(S, os.path.join(HERE, "slides"))
print(f"rendered {len(paths)} slides")
if WARNINGS:
    print("LAYOUT WARNINGS:")
    for w in WARNINGS:
        print("  -", w)
else:
    print("no layout warnings")

# ── pptx ────────────────────────────────────────────────────────────────────
from pptx import Presentation
from pptx.util import Inches

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]
for p, note in zip(paths, N):
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(p, 0, 0, width=prs.slide_width, height=prs.slide_height)
    slide.notes_slide.notes_text_frame.text = note
prs.save(os.path.join(HERE, "Together_Defense.pptx"))
print("pptx saved")

# ── pdf preview ─────────────────────────────────────────────────────────────
from PIL import Image
pages = [Image.open(p).convert("RGB").resize((1440, 810), Image.LANCZOS) for p in paths]
pages[0].save(os.path.join(HERE, "Together_Defense_preview.pdf"), save_all=True,
              append_images=pages[1:], resolution=110.0, quality=88)
print("pdf preview saved")

# ── contact sheets for inspection ───────────────────────────────────────────
COLS, ROWS = 4, 3
tw, th = 460, 259
for s in range((len(paths) + COLS * ROWS - 1) // (COLS * ROWS)):
    sheet = Image.new("RGB", (COLS * tw + 50, ROWS * th + 50), "#dddddd")
    for k in range(COLS * ROWS):
        idx = s * COLS * ROWS + k
        if idx >= len(paths):
            break
        im = Image.open(paths[idx]).resize((tw - 10, th - 10), Image.LANCZOS)
        r, c = divmod(k, COLS)
        sheet.paste(im, (10 + c * tw, 10 + r * th))
    sheet.save(os.path.join(HERE, f"contact_{s+1}.png"))
print("contact sheets saved")
